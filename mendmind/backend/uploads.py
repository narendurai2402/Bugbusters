"""Extract bounded text context from common student-uploaded files."""
from io import BytesIO
from pathlib import Path

import numpy as np
import pymupdf
from docx import Document
from PIL import Image
from pptx import Presentation
from rapidocr_onnxruntime import RapidOCR

MAX_EXTRACTED_CHARS = 12000
MAX_IMAGE_PIXELS = 24000000
MAX_PDF_PAGES = 30
MAX_SCANNED_PDF_PAGES = 5
_ocr_engine = None


def _ocr_image(image_bytes: bytes) -> str:
    global _ocr_engine
    if _ocr_engine is None:
        _ocr_engine = RapidOCR()

    with Image.open(BytesIO(image_bytes)) as source:
        if source.width * source.height > MAX_IMAGE_PIXELS:
            raise ValueError("Images must be 24 megapixels or smaller.")
        image = np.asarray(source.convert("RGB"))

    result = _ocr_engine(image)

    # rapidocr 3.x returns a RapidOCROutput object with a `txts` field.
    # Keep a fallback for older tuple/list-style OCR results.
    if hasattr(result, "txts"):
        texts = result.txts or []
    elif isinstance(result, tuple):
        results = result[0] if result else []
        texts = [
            item[1]
            for item in results
            if isinstance(item, (list, tuple)) and len(item) > 1
        ]
    else:
        texts = []

    return "\n".join(str(text).strip() for text in texts if str(text).strip())


def _extract_pdf(file_bytes: bytes) -> str:
    document = pymupdf.open(stream=file_bytes, filetype="pdf")
    if document.page_count > MAX_PDF_PAGES:
        raise ValueError(f"PDFs are limited to {MAX_PDF_PAGES} pages per upload.")

    page_text = [page.get_text("text").strip() for page in document]
    extracted = "\n\n".join(text for text in page_text if text)
    if extracted:
        return extracted

    ocr_text = []
    for page in document[:MAX_SCANNED_PDF_PAGES]:
        image_bytes = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).tobytes("png")
        text = _ocr_image(image_bytes)
        if text:
            ocr_text.append(text)
    return "\n\n".join(ocr_text)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(file_bytes: bytes) -> str:
    presentation = Presentation(BytesIO(file_bytes))
    if len(presentation.slides) > 100:
        raise ValueError("Presentations are limited to 100 slides per upload.")

    slides = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            slides.append(f"Slide {slide_number}\n" + "\n".join(parts))
    return "\n\n".join(slides)


def extract_attachment_text(filename: str, file_bytes: bytes) -> str:
    """Extract text from a supported file; OCR images and scanned PDF pages."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        text = _extract_pdf(file_bytes)
    elif extension == ".docx":
        text = _extract_docx(file_bytes)
    elif extension == ".pptx":
        text = _extract_pptx(file_bytes)
    elif extension in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        text = _ocr_image(file_bytes)
    elif extension in {
        ".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml",
        ".html", ".css", ".js", ".ts", ".py", ".sql", ".log",
    }:
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(f"{filename} is not valid UTF-8 text.") from exc
    else:
        raise ValueError(f"Unsupported file type for {filename}. Use PDF, DOCX, PPTX, an image, or a text/code file.")

    text = text.strip()
    if not text:
        raise ValueError(f"No readable text was found in {filename}.")
    return text[:MAX_EXTRACTED_CHARS]